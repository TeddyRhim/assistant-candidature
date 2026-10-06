from __future__ import annotations

import re
import unicodedata
import uuid
import zipfile
from collections.abc import Mapping
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from xml.etree.ElementTree import ParseError

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from pypdf.generic import ArrayObject, IndirectObject
from sqlalchemy import Engine, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from src.config import MAX_RESUME_SIZE_BYTES, get_data_dir
from src.models import ResumeVersion

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
SPACED_LETTER_RUN = re.compile(
    r"(?<!\S)(?:[^\W\d_]{1,2}[ \t]+)?(?:[^\W\d_][ \t]+){3,}"
    r"[^\W\d_](?=$|[\s,.;:!?])",
    re.UNICODE,
)
INVISIBLE_TEXT_CHARACTERS = re.compile("[\u00ad\u200b-\u200f\u2060\ufeff]")
TEXT_LIGATURES = str.maketrans(
    {
        "\ufb00": "ff",
        "\ufb01": "fi",
        "\ufb02": "fl",
        "\ufb03": "ffi",
        "\ufb04": "ffl",
    }
)


class ResumeImportError(ValueError):
    """A resume is unsupported, malformed, or cannot be extracted."""


def extract_resume_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeImportError("Format non pris en charge : sélectionnez un fichier PDF ou DOCX.")
    if not content:
        raise ResumeImportError("Le fichier sélectionné est vide.")
    if len(content) > MAX_RESUME_SIZE_BYTES:
        raise ResumeImportError("Le fichier dépasse la taille maximale de 10 Mo.")

    try:
        if extension == ".pdf":
            text = _extract_pdf_text(content)
        else:
            text = _extract_docx_text(content)
    except ResumeImportError:
        raise
    except (
        OSError,
        ValueError,
        KeyError,
        zipfile.BadZipFile,
        PackageNotFoundError,
        ParseError,
        PdfReadError,
    ) as error:
        raise ResumeImportError(
            "Le fichier ne peut pas être lu. Vérifiez qu'il s'agit d'un PDF ou DOCX valide."
        ) from error

    cleaned_text = normalize_resume_text(text)
    if not cleaned_text:
        if extension == ".pdf":
            raise ResumeImportError(
                "Aucun texte n'a été extrait. Le PDF est peut-être scanné ; l'OCR n'est pas "
                "pris en charge. Essayez un PDF textuel ou un fichier DOCX."
            )
        raise ResumeImportError("Aucun texte exploitable n'a été trouvé dans le document.")
    return cleaned_text


def _extract_pdf_text(content: bytes) -> str:
    if not content.startswith(b"%PDF-"):
        raise ResumeImportError("Le fichier sélectionné n'a pas la signature d'un PDF.")
    reader = PdfReader(BytesIO(content), strict=True)
    if reader.is_encrypted:
        raise ResumeImportError("Les PDF protégés par mot de passe ne sont pas pris en charge.")
    tagged_text = _extract_tagged_pdf_text(reader)
    layout_text = "\n".join(
        page.extract_text(extraction_mode="layout") or ""
        for page in reader.pages
        if page.get("/Contents") is not None
    )
    normalized_tagged = normalize_resume_text(tagged_text)
    normalized_layout = normalize_resume_text(layout_text)
    if normalized_tagged and (
        not normalized_layout or len(normalized_tagged) >= len(normalized_layout) * 0.7
    ):
        return tagged_text
    return layout_text


def _extract_tagged_pdf_text(reader: PdfReader) -> str:
    root = reader.trailer.get("/Root")
    if isinstance(root, IndirectObject):
        root = root.get_object()
    if not isinstance(root, Mapping):
        return ""
    structure_tree = root.get("/StructTreeRoot")
    if isinstance(structure_tree, IndirectObject):
        structure_tree = structure_tree.get_object()
    if not isinstance(structure_tree, Mapping):
        return ""

    extracted: list[str] = []

    def visit(item: object) -> None:
        if isinstance(item, IndirectObject):
            visit(item.get_object())
        elif isinstance(item, Mapping):
            actual_text = item.get("/E")
            if actual_text is not None:
                value = str(actual_text).strip()
                if value:
                    extracted.append(value)
            else:
                visit(item.get("/K"))
        elif isinstance(item, (ArrayObject, list, tuple)):
            for child in item:
                visit(child)

    visit(structure_tree.get("/K"))
    return "\n".join(extracted)


def normalize_resume_text(text: str) -> str:
    normalized = unicodedata.normalize("NFC", text).translate(TEXT_LIGATURES)
    normalized = INVISIBLE_TEXT_CHARACTERS.sub("", normalized).replace("\u00a0", " ")

    def join_spaced_letters(match: re.Match[str]) -> str:
        return re.sub(r"[ \t]+", "", match.group(0))

    lines = []
    for line in normalized.splitlines():
        line = SPACED_LETTER_RUN.sub(join_spaced_letters, line)
        line = re.sub(r"[ \t]{2,}", " ", line).strip()
        lines.append(line)
    normalized = re.sub(r"\n{3,}", "\n\n", "\n".join(lines))
    return normalized.strip()


def _extract_docx_text(content: bytes) -> str:
    if not content.startswith(b"PK"):
        raise ResumeImportError("Le fichier sélectionné n'a pas la signature d'un DOCX.")
    with zipfile.ZipFile(BytesIO(content)) as archive:
        total_uncompressed_size = sum(item.file_size for item in archive.infolist())
        if total_uncompressed_size > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise ResumeImportError(
                "Le contenu DOCX décompressé dépasse la limite de sécurité autorisée."
            )
        required_entries = {"[Content_Types].xml", "word/document.xml"}
        if not required_entries.issubset(archive.namelist()):
            raise ResumeImportError("Le fichier sélectionné n'est pas un document DOCX valide.")

    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_text = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *table_text])


def save_resume_version(
    engine: Engine,
    filename: str,
    content: bytes,
    reviewed_text: str,
) -> ResumeVersion:
    extract_resume_text(filename, content)
    if not reviewed_text.strip():
        raise ResumeImportError("Le texte vérifié du CV ne peut pas être vide.")
    if len(filename) > 255:
        filename = filename[-255:]

    resume_id = uuid.uuid4().hex
    extension = Path(filename).suffix.lower()
    stored_filename = f"{resume_id}{extension}"
    resume_directory = get_data_dir() / "resumes"
    resume_directory.mkdir(parents=True, exist_ok=True)
    stored_path = resume_directory / stored_filename
    stored_path.write_bytes(content)

    resume = ResumeVersion(
        id=resume_id,
        original_filename=Path(filename).name,
        stored_filename=stored_filename,
        reviewed_text=reviewed_text.strip(),
        created_at=datetime.now(UTC).isoformat(),
    )
    try:
        with Session(engine, expire_on_commit=False) as session, session.begin():
            session.add(resume)
    except SQLAlchemyError:
        stored_path.unlink(missing_ok=True)
        raise
    return resume


def list_resume_versions(engine: Engine) -> list[ResumeVersion]:
    with Session(engine) as session:
        return list(
            session.scalars(select(ResumeVersion).order_by(ResumeVersion.created_at.desc()))
        )
