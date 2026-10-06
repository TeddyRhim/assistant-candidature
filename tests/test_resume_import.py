from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfWriter

from src import db
from src.config import DATA_DIRECTORY_ENV
from src.services.resume_import import (
    ResumeImportError,
    _extract_tagged_pdf_text,
    extract_resume_text,
    list_resume_versions,
    normalize_resume_text,
    save_resume_version,
)


def make_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_blank_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_extracts_paragraphs_and_tables_from_docx() -> None:
    document = Document()
    document.add_paragraph("Développeur backend")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Symfony"
    buffer = BytesIO()
    document.save(buffer)

    text = extract_resume_text("cv.docx", buffer.getvalue())

    assert text == "Développeur backend\nSymfony"


def test_normalizes_spaced_letters_unicode_and_invisible_characters() -> None:
    text = "S y m f o n y · De\u0301 v e l o p p e u r\u00ad  PHP · o\ufb03ce"

    normalized = normalize_resume_text(text)

    assert normalized == "Symfony · Développeur PHP · office"


def test_extracts_semantic_text_from_tagged_pdf_in_document_order() -> None:
    reader = SimpleNamespace(
        trailer={
            "/Root": {
                "/StructTreeRoot": {
                    "/K": {
                        "/S": "/Document",
                        "/K": [
                            {"/S": "/P", "/E": "TEDDY RHIM"},
                            {
                                "/S": "/Sect",
                                "/K": [
                                    {"/S": "/H2", "/E": "Compétences"},
                                    {"/S": "/P", "/E": "Symfony"},
                                ],
                            },
                        ],
                    }
                }
            }
        }
    )

    assert _extract_tagged_pdf_text(reader) == "TEDDY RHIM\nCompétences\nSymfony"


def test_extracts_pdf_text_in_layout_mode_and_repairs_character_spacing() -> None:
    from reportlab.pdfgen.canvas import Canvas

    buffer = BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(72, 720, "S y m f o n y")
    canvas.save()

    text = extract_resume_text("cv.pdf", buffer.getvalue())

    assert text == "Symfony"


def test_rejects_unsupported_type_and_wrong_file_signature() -> None:
    with pytest.raises(ResumeImportError, match="Format non pris en charge"):
        extract_resume_text("cv.txt", b"not a supported file")
    with pytest.raises(ResumeImportError, match="signature"):
        extract_resume_text("cv.pdf", b"not a pdf")
    with pytest.raises(ResumeImportError, match="signature"):
        extract_resume_text("cv.docx", b"not a docx")


def test_rejects_malformed_docx_and_scanned_or_empty_pdf() -> None:
    with pytest.raises(ResumeImportError, match="PDF ou DOCX valide"):
        extract_resume_text("cv.docx", b"PKnot a zip archive")
    with pytest.raises(ResumeImportError, match="scanné"):
        extract_resume_text("cv.pdf", make_blank_pdf_bytes())


def test_rejects_password_protected_pdf() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.encrypt("secret")
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(ResumeImportError, match="mot de passe"):
        extract_resume_text("cv.pdf", buffer.getvalue())


def test_rejects_oversized_and_empty_files() -> None:
    with pytest.raises(ResumeImportError, match="vide"):
        extract_resume_text("cv.docx", b"")
    with pytest.raises(ResumeImportError, match="10 Mo"):
        extract_resume_text("cv.pdf", b"%PDF-" + b"x" * (10 * 1024 * 1024))


def test_saves_reviewed_resume_and_lists_persisted_version(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv(DATA_DIRECTORY_ENV, str(tmp_path / "personal-data"))
    engine = db.create_database_engine(tmp_path / "profile.sqlite3")
    db.initialize_database(engine)
    content = make_docx_bytes("Texte extrait")

    resume = save_resume_version(
        engine,
        "cv-reference.docx",
        content,
        "Texte corrigé par l'utilisateur",
    )

    assert resume.original_filename == "cv-reference.docx"
    assert resume.reviewed_text == "Texte corrigé par l'utilisateur"
    stored_path = tmp_path / "personal-data" / "resumes" / resume.stored_filename
    assert stored_path.read_bytes() == content
    assert [item.id for item in list_resume_versions(engine)] == [resume.id]
    engine.dispose()


def test_rejects_blank_reviewed_text(tmp_path) -> None:
    engine = db.create_database_engine(tmp_path / "profile.sqlite3")
    db.initialize_database(engine)

    with pytest.raises(ResumeImportError, match="ne peut pas être vide"):
        save_resume_version(engine, "cv.docx", make_docx_bytes("Texte"), "  ")
    engine.dispose()
